from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)

# ── Helpers ───────────────────────────────────────────────────────────────────
def add_heading(doc, text, level, color=None):
    para = doc.add_heading(text, level=level)
    for run in para.runs:
        sizes = {1: 22, 2: 16, 3: 13, 4: 11}
        run.font.size = Pt(sizes.get(level, 12))
        run.font.bold = True
        run.font.color.rgb = RGBColor(*color) if color else RGBColor(0x1a, 0x18, 0x10)
    para.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
    para.paragraph_format.space_after  = Pt(4)
    return para

def add_para(doc, text, bold=False, italic=False, size=10.5, color=None, indent=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_after  = Pt(4)
    para.paragraph_format.space_before = Pt(0)
    if indent:
        para.paragraph_format.left_indent = Cm(indent)
    run = para.add_run(text)
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return para

def add_bullet(doc, text, indent_level=0):
    para = doc.add_paragraph(style='List Bullet')
    para.paragraph_format.space_after  = Pt(2)
    para.paragraph_format.space_before = Pt(0)
    if indent_level:
        para.paragraph_format.left_indent = Cm(1 + indent_level * 0.5)
    run = para.add_run(text)
    run.font.size = Pt(10.5)
    return para

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  '2C5F4A')
        tcPr.append(shd)

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        fill = 'F5F2EA' if r_idx % 2 == 0 else 'FFFFFF'
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(cell_text)
            run.font.size = Pt(10)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  fill)
            tcPr.append(shd)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table

def add_note_box(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent  = Cm(0.5)
    para.paragraph_format.right_indent = Cm(0.5)
    para.paragraph_format.space_after  = Pt(8)
    run = para.add_run('💡  ' + text)
    run.font.size   = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x2C, 0x5F, 0x4A)

# ══════════════════════════════════════════════════════════════════════════════
# TRANG BÌA
# ══════════════════════════════════════════════════════════════════════════════
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_para.paragraph_format.space_before = Pt(40)
run = title_para.add_run('Công Cụ Hỗ Trợ Quyết Định')
run.font.size  = Pt(28)
run.font.bold  = True
run.font.color.rgb = RGBColor(0x2C, 0x5F, 0x4A)

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = sub_para.add_run('Decision Helper — Hướng dẫn sử dụng')
run2.font.size   = Pt(14)
run2.font.italic = True
run2.font.color.rgb = RGBColor(0x6B, 0x68, 0x59)

doc.add_paragraph()
add_para(doc, '─' * 80, size=8, color=(0xDD, 0xD9, 0xCF))
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# TỔNG QUAN
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Tổng quan', 1)
add_para(doc,
    'Công cụ Hỗ trợ Quyết định (Decision Helper) là trợ lý giao dịch dựa trên quy tắc, '
    'giúp phân tích cổ phiếu Việt Nam và đưa ra khuyến nghị: nên làm gì, khi nào, '
    'và lý do tại sao. Hệ thống sử dụng quy trình nhiều bước:')
add_para(doc,
    'Xác định Xu hướng Thị trường  →  Tạo Tín hiệu  →  Kiểm tra Rủi ro & Kỷ luật  '
    '→  Tính Khối lượng Vị thế  →  Khuyến nghị Cuối cùng',
    bold=True, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Menu này gồm bốn màn hình:')
for item in [
    '1.  Bảng điều khiển Quyết định (Decision Console) — phân tích chi tiết từng mã',
    '2.  Bảng Quyết định Danh mục (Decision Board)     — bảng tổng hợp toàn bộ danh sách theo dõi',
    '3.  Bản đồ nhiệt Danh mục (Watchlist Heat Map)   — lưới thẻ trực quan để quét nhanh',
    '4.  Bảng Kỷ luật (Discipline Dashboard)           — theo dõi tuân thủ quy tắc giao dịch',
]:
    add_para(doc, item, indent=0.5)

# ══════════════════════════════════════════════════════════════════════════════
# 1. BẢNG ĐIỀU KHIỂN QUYẾT ĐỊNH
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '1.  Bảng điều khiển Quyết định (Decision Console)', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Đường dẫn: /decision-console', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'Đây là màn hình cốt lõi. Nhập một mã cổ phiếu và nhận phân tích đầy đủ '
    'cho một ngày giao dịch cụ thể.')

add_heading(doc, 'Cách sử dụng', 2)
add_bullet(doc, 'Nhập mã cổ phiếu vào ô tìm kiếm (ví dụ: VCB, FPT, VNM). Hệ thống tự động chuyển sang chữ hoa.')
add_bullet(doc, 'Chọn ngày phân tích — mặc định là ngày hôm nay. Có thể phân tích bất kỳ ngày nào trong quá khứ.')
add_bullet(doc, 'Nhấn nút Phân tích hoặc nhấn Enter để chạy toàn bộ quy trình quyết định.')

add_heading(doc, 'Đọc kết quả phân tích', 2)
add_table(doc,
    ['Thành phần', 'Ý nghĩa'],
    [
        ['Nhãn Xu hướng (Regime)',   'Trạng thái thị trường: TREND UP — tăng (xanh), TREND DOWN — giảm (đỏ), RANGE — đi ngang (vàng), UNDEFINED — chưa xác định (xám)'],
        ['Nhãn Hành động (Action)',  'Khuyến nghị của hệ thống: MUA (BUY), BÁN (SELL), BÁN KHỐNG (SELL SHORT), ĐÓI MUA (COVER), GIỮ (HOLD), CHỜ (WAIT), ĐẢO CHIỀU (REVERSE)'],
        ['Độ tin cậy (%)',           'Mức độ đồng thuận của tín hiệu qua tất cả các bước trong quy trình (0–100%)'],
        ['Giá vào / Cắt lỗ / Mục tiêu / Khối lượng', 'Các mức giá giao dịch được đề xuất. Cắt lỗ và Mục tiêu bắt buộc phải có trước khi lưu kế hoạch'],
        ['Thanh R:R',                'Tỷ lệ lợi nhuận/rủi ro. Xanh ≥ 1.5:1, vàng < 1.5:1. Cảnh báo hiện ra khi tỷ lệ dưới 1:1'],
        ['Mô hình giá (Pattern)',    'Các mô hình biểu đồ được phát hiện (Đầu & Vai, Đáy Đôi,...) kèm thanh độ tin cậy'],
        ['Cảnh báo Kỷ luật',        'BLOCK = vi phạm quy tắc, không cho phép vào lệnh. WARN = cảnh báo, vẫn có thể giao dịch nhưng cần thận trọng'],
        ['Chuỗi lý do (Reason Chain)','Nhấn để mở rộng toàn bộ logic từng bước: xu hướng → tín hiệu → kỷ luật → khối lượng'],
    ],
    col_widths=[4.5, 11]
)

add_heading(doc, 'Lưu Kế hoạch Giao dịch', 2)
add_para(doc,
    'Khi hành động là MUA (BUY) hoặc BÁN KHỐNG (SELL SHORT), nút Lưu Kế hoạch Giao dịch '
    'sẽ xuất hiện. Nhấn vào để mở biểu mẫu:')
add_bullet(doc, 'Điều chỉnh giá vào lệnh, cắt lỗ và mục tiêu nếu cần.')
add_bullet(doc, 'Chọn tài khoản nhật ký.')
add_bullet(doc, 'Thêm ghi chú hoặc lý do thiết lập giao dịch.')
add_bullet(doc, 'Nhấn Lưu Kế hoạch — hệ thống tạo bản ghi giao dịch trong Nhật ký và tự động '
                'tạo mục nhật ký với ngữ cảnh tín hiệu và chuỗi lý do đầy đủ.')

add_heading(doc, 'Cài đặt Rủi ro  (biểu tượng ⚙)', 2)
add_para(doc,
    'Nhấn biểu tượng cài đặt (góc trên bên phải thanh công cụ) để cấu hình bộ tính khối lượng:')
add_table(doc,
    ['Cài đặt', 'Mô tả'],
    [
        ['Vốn tài khoản (Account Equity)',      'Tổng vốn đầu tư. Dùng để tính khối lượng lệnh theo giá trị tiền tệ.'],
        ['Rủi ro mỗi lệnh (%)',                 'Tỷ lệ % vốn tối đa chấp nhận rủi ro cho một lệnh (khuyến nghị: 0.5–2%).'],
        ['Số vị thế tối đa',                    'Số lệnh mở đồng thời tối đa mà hệ thống cho phép.'],
        ['Tương quan tối đa',                   'Hệ số tương quan tối đa được phép giữa các vị thế đang mở (0–1). Ngăn chặn tập trung rủi ro vào cùng một ngành.'],
        ['Hệ số giảm khối lượng',               'Hệ số nhân khi đang trong chuỗi thua lỗ (ví dụ: 0.5 = giảm còn một nửa khối lượng).'],
        ['Ngưỡng chuỗi thua',                   'Số lệnh thua liên tiếp để kích hoạt giảm khối lượng.'],
        ['ID tài khoản Nhật ký',                'Tùy chọn. Liên kết tính khối lượng với một tài khoản nhật ký cụ thể để theo dõi P&L.'],
    ],
    col_widths=[4.5, 11]
)

# ══════════════════════════════════════════════════════════════════════════════
# 2. BẢNG QUYẾT ĐỊNH DANH MỤC
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '2.  Bảng Quyết định Danh mục (Decision Board)', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Đường dẫn: /decision-board', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'Bảng tổng hợp nhiều mã cổ phiếu, chạy phân tích đầy đủ song song cho toàn bộ danh sách '
    'theo dõi. Phù hợp nhất để xem lại buổi sáng hàng ngày trước khi vào lệnh.')

add_heading(doc, 'Cách sử dụng', 2)
add_bullet(doc, 'Danh sách theo dõi được tải tự động từ module Danh mục đầu tư (Portfolio).')
add_bullet(doc, 'Chọn ngày phân tích (mặc định hôm nay).')
add_bullet(doc, 'Nhấn Phân tích tất cả (N) — hệ thống chạy tất cả mã song song. Mỗi hàng hiển thị vòng quay trong khi xử lý.')
add_bullet(doc, 'Nhấn biểu tượng ↻ ở bất kỳ hàng nào để chạy lại riêng mã đó.')

add_heading(doc, 'Đọc bảng kết quả', 2)
add_table(doc,
    ['Cột', 'Ý nghĩa'],
    [
        ['Mã CK (Ticker)',      'Nhấn để chuyển sang Bảng điều khiển Quyết định với phân tích chi tiết của mã đó.'],
        ['Xu hướng (Regime)',   'Trạng thái thị trường, được tô màu (xanh / đỏ / vàng / xám).'],
        ['Hành động (Action)',  'Nhãn khuyến nghị cho mã cổ phiếu đó.'],
        ['Độ tin cậy (%)',      'Điểm độ tin cậy của tín hiệu.'],
        ['Vào / Cắt / Mục tiêu / Khối lượng', 'Các mức giá giao dịch được đề xuất.'],
        ['↻ (làm mới)',         'Chạy lại phân tích riêng cho mã đó.'],
        ['📖 (sổ)',             'Lưu kế hoạch giao dịch (chỉ hiện với tín hiệu MUA hoặc BÁN KHỐNG).'],
    ],
    col_widths=[4, 11.5]
)

add_heading(doc, 'Thẻ tổng hợp', 2)
add_para(doc,
    'Sau khi phân tích xong, các thẻ màu xuất hiện phía trên bảng, đếm số mã theo từng '
    'trạng thái hành động (ví dụ: MUA 3 · GIỮ 5 · CHỜ 4). Cho cái nhìn tổng thể nhanh về '
    'tình trạng toàn thị trường.')

add_heading(doc, 'Sắp xếp theo cột', 2)
add_para(doc,
    'Nhấn vào tiêu đề bất kỳ để sắp xếp tăng/giảm dần — hữu ích để tìm tín hiệu MUA '
    'có độ tin cậy cao nhất hoặc nhóm tất cả mã cùng xu hướng lại với nhau.')

# ══════════════════════════════════════════════════════════════════════════════
# 3. BẢN ĐỒ NHIỆT DANH MỤC
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '3.  Bản đồ nhiệt Danh mục (Watchlist Heat Map)', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Đường dẫn: /watchlist-heatmap', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'Lưới thẻ trực quan hiển thị tất cả mã trong danh sách theo dõi, tô màu theo xu hướng '
    'thị trường và tín hiệu hành động. Thiết kế để quét thị trường nhanh chóng trước khi '
    'đi vào chi tiết.')

add_heading(doc, 'Cách sử dụng', 2)
add_bullet(doc, 'Chọn ngày (mặc định hôm nay).')
add_bullet(doc, 'Nhấn Quét ngay (Scan Now) — hệ thống quét tất cả mã và hiển thị kết quả dạng thẻ.')
add_bullet(doc, 'Nhấn vào bất kỳ thẻ nào để mở Bảng điều khiển Quyết định cho mã đó.')

add_heading(doc, 'Đọc thẻ kết quả', 2)
add_para(doc, 'Mỗi thẻ hiển thị:')
add_bullet(doc, 'Tên mã cổ phiếu — chữ lớn, màu theo xu hướng thị trường.')
add_bullet(doc, 'Nhãn hành động — MUA, BÁN KHỐNG, BÁN, GIỮ, CHỜ, v.v.')
add_bullet(doc, 'Chấm xu hướng + nhãn — trạng thái thị trường kèm chấm màu chỉ thị.')
add_bullet(doc, 'Loại tín hiệu — ví dụ: MA_CROSS (cắt MA), BREAKOUT (phá vỡ), STOCH_CROSS (cắt Stochastic).')
add_bullet(doc, 'Thanh độ tin cậy — xanh ≥ 70%, vàng 50–69%, đỏ < 50%.')
add_bullet(doc, 'Mức giá — Vào lệnh, Cắt lỗ, Mục tiêu (chỉ hiện khi có tín hiệu vào lệnh).')
add_para(doc, 'Thẻ có tín hiệu vào lệnh (MUA hoặc BÁN KHỐNG) được viền màu nổi bật để dễ nhận ra.')

add_heading(doc, 'Lọc & Sắp xếp', 2)
add_table(doc,
    ['Điều khiển', 'Tùy chọn'],
    [
        ['Thanh lọc', 'Tất cả  |  Có thể vào lệnh  |  MUA  |  BÁN KHỐNG  |  GIỮ  |  CHỜ'],
        ['Sắp xếp',  'Độ tin cậy (mặc định)  |  Mã CK  |  Hành động  |  Xu hướng'],
    ],
    col_widths=[4, 11.5]
)
add_note_box(doc, 'Dùng bộ lọc "Có thể vào lệnh" để thu hẹp lưới chỉ còn những mã có tín hiệu vào lệnh ngay lập tức.')

# ══════════════════════════════════════════════════════════════════════════════
# 4. BẢNG KỶ LUẬT
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '4.  Bảng Kỷ luật (Discipline Dashboard)', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Đường dẫn: /discipline', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'Theo dõi kỷ luật giao dịch dựa trên các quy tắc được chuyển thể từ cuốn sách '
    'Market Wizards của Jack Schwager. Hệ thống phân tích lịch sử giao dịch gần đây '
    'và cảnh báo khi bạn đang vi phạm quy tắc quản lý rủi ro.')

add_heading(doc, 'Thẻ thống kê (hàng trên cùng)', 2)
add_table(doc,
    ['Thẻ', 'Nội dung hiển thị'],
    [
        ['Chuỗi thua lỗ',         'Số lệnh thua liên tiếp gần đây.'],
        ['Chuỗi đang hoạt động',  'CÓ / Không — có đang trong chuỗi thua vượt ngưỡng hay không. Khi hiện CÓ, khối lượng lệnh được tự động giảm.'],
        ['Drawdown (%)',           'Mức suy giảm hiện tại so với đỉnh vốn. Drawdown ≥ 20% kích hoạt quy tắc bắt buộc xem xét lại chiến lược.'],
        ['Quy tắc bị vi phạm',    'Số quy tắc Schwager đang bị vi phạm trên tổng số quy tắc trong bộ sách.'],
    ],
    col_widths=[4, 11.5]
)
add_para(doc, 'Thẻ có viền vàng cho biết điều kiện cảnh báo cần xử lý ngay.')

add_heading(doc, '10 Hành động gần nhất', 2)
add_para(doc, 'Chuỗi 10 hành động giao dịch gần nhất, tô màu:')
add_bullet(doc, 'Xanh  = lệnh vào (MUA / BÁN KHỐNG)')
add_bullet(doc, 'Đỏ    = lệnh thoát (BÁN / ĐÓI MUA)')
add_bullet(doc, 'Xám   = trung tính (GIỮ / CHỜ)')

add_heading(doc, 'Vi phạm đang hoạt động', 2)
add_table(doc,
    ['Mức độ', 'Ý nghĩa'],
    [
        ['🔴  BLOCK (Chặn)', 'Vi phạm cứng. Bảng điều khiển Quyết định sẽ tự động từ chối tín hiệu MUA / BÁN KHỐNG khi vi phạm này còn tồn tại.'],
        ['🟡  WARN (Cảnh báo)', 'Vi phạm mềm. Vẫn có thể giao dịch nhưng hệ thống sẽ ghi nhận rủi ro trong chuỗi lý do.'],
    ],
    col_widths=[3.5, 12]
)

add_heading(doc, 'Bảng Quy tắc Schwager', 2)
add_para(doc,
    'Toàn bộ bộ quy tắc được hiển thị — quy tắc đang vi phạm (tô màu nổi bật) và quy tắc '
    'đang tuân thủ (mờ hơn). Mỗi quy tắc bao gồm:')
add_bullet(doc, 'Mã quy tắc và tham chiếu chương trong Market Wizards')
add_bullet(doc, 'Nội dung quy tắc')
add_bullet(doc, 'Thông báo chi tiết giải thích lý do vi phạm (khi đang bị kích hoạt)')

add_heading(doc, 'Khi nào nên kiểm tra màn hình này', 2)
add_bullet(doc, 'Trước khi quét thị trường — vi phạm BLOCK sẽ tự động chặn tín hiệu vào lệnh.')
add_bullet(doc, 'Sau một chuỗi thua lỗ — bảng xác nhận khối lượng giảm có đang được áp dụng không.')
add_bullet(doc, 'Sau khi drawdown lớn — quy tắc drawdown 20% kích hoạt yêu cầu xem xét lại bắt buộc.')

# ══════════════════════════════════════════════════════════════════════════════
# QUY TRÌNH KHUYẾN NGHỊ
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Quy trình Sử dụng Hàng ngày', 1)
steps = [
    ('Bước 1', 'Bảng Kỷ luật',          'Kiểm tra xem có vi phạm BLOCK nào không. Nếu có, giải quyết vấn đề trước khi giao dịch.'),
    ('Bước 2', 'Bản đồ nhiệt Danh mục', 'Quét thị trường ngày hôm nay. Dùng bộ lọc "Có thể vào lệnh" để tìm tín hiệu vào lệnh.'),
    ('Bước 3', 'Bảng điều khiển QĐ',   'Nhấn vào từng thẻ tiềm năng. Xem xét xu hướng, tín hiệu, chuỗi lý do và tỷ lệ R:R.'),
    ('Bước 4', 'Lưu Kế hoạch',          'Nếu tín hiệu đạt tiêu chí (độ tin cậy ≥ 60%, R:R ≥ 1.5:1), lưu kế hoạch giao dịch.'),
    ('Bước 5', 'Bảng Quyết định',       'Kiểm tra lại bảng tổng hợp cuối ngày để theo dõi thay đổi xu hướng toàn danh mục.'),
]
add_table(doc,
    ['Bước', 'Màn hình', 'Hành động'],
    steps,
    col_widths=[2, 4.5, 9]
)

# ══════════════════════════════════════════════════════════════════════════════
# MẸO VÀ THỰC HÀNH TỐT
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Mẹo & Thực hành Tốt', 1)
tips = [
    ('Phân tích lịch sử',     'Dùng bất kỳ ngày nào trong quá khứ ở ô ngày để kiểm tra tín hiệu trên dữ liệu lịch sử.'),
    ('Ngưỡng độ tin cậy',     'Tín hiệu dưới 60% là độ tin cậy thấp — hãy chờ thiết lập mạnh hơn.'),
    ('Tỷ lệ R:R',             'Chỉ vào lệnh khi R:R ≥ 1.5:1. Hệ thống cảnh báo khi tỷ lệ dưới 1:1.'),
    ('GIỮ vs CHỜ',            'GIỮ = bạn đang có vị thế mở, tiếp tục nắm giữ. CHỜ = chưa có vị thế và chưa có tín hiệu vào lệnh.'),
    ('Xác nhận mô hình giá',  'Mô hình biểu đồ phát hiện đồng thuận với hành động (ví dụ: Đáy Đôi + MUA) tăng đáng kể mức độ tin tưởng.'),
    ('Tính khối lượng',       'Nhập chính xác Vốn tài khoản và % Rủi ro trong Cài đặt Rủi ro — con số Khối lượng phụ thuộc vào hai giá trị này.'),
    ('Chuỗi thua lỗ',         'Nếu "Chuỗi đang hoạt động" hiện CÓ, hệ thống đã tự động giảm khối lượng — không cần can thiệp thủ công.'),
]
add_table(doc,
    ['Mẹo', 'Chi tiết'],
    tips,
    col_widths=[4.5, 11]
)

# ══════════════════════════════════════════════════════════════════════════════
# LƯU FILE
# ══════════════════════════════════════════════════════════════════════════════
out_path = '/home/tuanph/all_projects/stockpredict/Huong_Dan_Su_Dung_Cong_Cu_Ho_Tro_Quyet_Dinh.docx'
doc.save(out_path)
print(f'Saved: {out_path}')
