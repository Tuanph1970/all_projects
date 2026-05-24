from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)

# ── Styles helpers ────────────────────────────────────────────────────────────
def set_heading(para, text, level, color=None):
    para.clear()
    run = para.add_run(text)
    sizes = {1: 22, 2: 16, 3: 13, 4: 11}
    run.font.size = Pt(sizes.get(level, 12))
    run.font.bold = True
    if color:
        run.font.color.rgb = RGBColor(*color)
    else:
        run.font.color.rgb = RGBColor(0x1a, 0x18, 0x10)

def add_heading(doc, text, level, color=None):
    style = f'Heading {level}'
    para = doc.add_heading(text, level=level)
    for run in para.runs:
        sizes = {1: 22, 2: 16, 3: 13, 4: 11}
        run.font.size = Pt(sizes.get(level, 12))
        run.font.bold = True
        run.font.color.rgb = RGBColor(*color) if color else RGBColor(0x1a, 0x18, 0x10)
    para.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
    para.paragraph_format.space_after  = Pt(4)
    return para

def add_para(doc, text, bold=False, italic=False, size=10.5, color=None, indent=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_after  = Pt(4)
    para.paragraph_format.space_before = Pt(0)
    if indent:
        para.paragraph_format.left_indent = Cm(indent)
    run = para.add_run(text)
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return para

def add_bullet(doc, text, indent_level=0):
    para = doc.add_paragraph(style='List Bullet')
    para.paragraph_format.space_after  = Pt(2)
    para.paragraph_format.space_before = Pt(0)
    if indent_level:
        para.paragraph_format.left_indent = Cm(1 + indent_level * 0.5)
    run = para.add_run(text)
    run.font.size = Pt(10.5)
    return para

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  '2C5F4A')
        tcPr.append(shd)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        fill = 'F5F2EA' if r_idx % 2 == 0 else 'FFFFFF'
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.paragraphs[0].clear()
            run = cell.paragraphs[0].add_run(cell_text)
            run.font.size = Pt(10)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'),   'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'),  fill)
            tcPr.append(shd)

    # Column widths
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table

def add_note_box(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent  = Cm(0.5)
    para.paragraph_format.right_indent = Cm(0.5)
    para.paragraph_format.space_after  = Pt(8)
    run = para.add_run('💡  ' + text)
    run.font.size   = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x2C, 0x5F, 0x4A)

# ══════════════════════════════════════════════════════════════════════════════
# COVER
# ══════════════════════════════════════════════════════════════════════════════
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_para.paragraph_format.space_before = Pt(40)
run = title_para.add_run('Decision Helper')
run.font.size  = Pt(32)
run.font.bold  = True
run.font.color.rgb = RGBColor(0x2C, 0x5F, 0x4A)

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = sub_para.add_run('Công cụ hỗ trợ quyết định — User Guide')
run2.font.size   = Pt(14)
run2.font.italic = True
run2.font.color.rgb = RGBColor(0x6B, 0x68, 0x59)

doc.add_paragraph()
add_para(doc, '─' * 80, size=8, color=(0xDD, 0xD9, 0xCF))

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# OVERVIEW
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Overview', 1)
add_para(doc,
    'The Decision Helper is a rules-based trading assistant that analyses Vietnamese stocks '
    'and tells you what to do, when, and why. It uses a multi-stage pipeline:')
add_para(doc,
    'Market Regime Detection  →  Signal Generation  →  Risk & Discipline Checks  '
    '→  Position Sizing  →  Final Recommendation',
    bold=True, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'There are four screens under this menu:')
for item in [
    '1.  Decision Console — deep single-ticker analysis',
    '2.  Decision Board   — watchlist-wide summary table',
    '3.  Watchlist Heat Map — visual card grid for quick scanning',
    '4.  Discipline Dashboard — trading-rules compliance monitor',
]:
    add_para(doc, item, indent=0.5)

# ══════════════════════════════════════════════════════════════════════════════
# 1. DECISION CONSOLE
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '1.  Decision Console', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Route: /decision-console', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'This is the core screen. Enter one ticker symbol and get a complete analysis '
    'for a specific trading date.')

add_heading(doc, 'How to use', 2)
add_bullet(doc, 'Type a ticker symbol in the input box (e.g. VCB, FPT, VNM). It auto-uppercases.')
add_bullet(doc, 'Select a date — defaults to today. You can analyse any past bar date.')
add_bullet(doc, 'Click Analyze or press Enter to run the full decision pipeline.')

add_heading(doc, 'Reading the result', 2)
add_table(doc,
    ['Element', 'What it means'],
    [
        ['Regime pill',    'Market state: TREND UP (green), TREND DOWN (red), RANGE (amber), UNDEFINED (grey)'],
        ['Action badge',   'Recommendation: BUY, SELL, SELL SHORT, COVER, HOLD, WAIT, or REVERSE'],
        ['Confidence %',   'How strongly the signal aligns across all pipeline stages (0–100%)'],
        ['Entry / Stop / Target / Shares', 'Suggested trade levels. Stop and Target are required before saving a plan'],
        ['R:R bar',        'Risk-to-reward ratio. Green ≥ 1.5:1, amber < 1.5:1. A warning appears below 1:1'],
        ['Pattern overlay','Detected chart patterns (e.g. Head & Shoulders, Double Top) with confidence bar'],
        ['Discipline warnings', 'BLOCK = rule violation preventing entry.  WARN = caution, trade allowed'],
        ['Reason chain',   'Click to expand the full step-by-step logic: regime → signal → discipline → sizing'],
    ],
    col_widths=[4.5, 11]
)

add_heading(doc, 'Saving a Trade Plan', 2)
add_para(doc,
    'When the action is BUY or SELL SHORT, a Save Trade Plan button appears. Click it to open '
    'the trade plan form:')
add_bullet(doc, 'Adjust entry, stop, and target prices if needed.')
add_bullet(doc, 'Select a journal account.')
add_bullet(doc, 'Add notes or setup reasoning.')
add_bullet(doc, 'Click Save Trade Plan — this creates a trade record in your Journal and '
                'auto-generates a diary entry with signal context and reason chain.')

add_heading(doc, 'Risk Preferences  (⚙ settings icon)', 2)
add_para(doc,
    'Click the gear icon (top-right of the toolbar) to configure the sizing engine:')
add_table(doc,
    ['Setting', 'Description'],
    [
        ['Account Equity',          'Your total capital. Used to calculate position sizes in currency terms.'],
        ['Risk per Trade %',        'Maximum % of equity to risk on one trade (recommended: 0.5–2%).'],
        ['Max Concurrent Positions','How many open trades the engine will allow simultaneously.'],
        ['Max Correlation',         'Maximum allowed correlation between open positions (0–1). Prevents over-exposure to the same sector.'],
        ['Loss Streak Cut',         'Position-size multiplier applied when on a losing streak (e.g. 0.5 = half size).'],
        ['Loss Streak Threshold',   'Number of consecutive losses that triggers the size reduction.'],
        ['Journal Account ID',      'Optional. Links position sizing to a specific journal account for P&L tracking.'],
    ],
    col_widths=[4.5, 11]
)

# ══════════════════════════════════════════════════════════════════════════════
# 2. DECISION BOARD
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '2.  Decision Board', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Route: /decision-board', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'A multi-symbol table that runs the full analysis across your entire watchlist in parallel. '
    'Best used for a daily morning review.')

add_heading(doc, 'How to use', 2)
add_bullet(doc, 'Your watchlist is loaded automatically from the Portfolio module.')
add_bullet(doc, 'Select a date (defaults to today).')
add_bullet(doc, 'Click Analyze All (N) — the engine runs all tickers in parallel. '
                'Each row shows a spinner while processing.')
add_bullet(doc, 'Use the ↻ icon on any row to re-run analysis for that single ticker.')

add_heading(doc, 'Reading the table', 2)
add_table(doc,
    ['Column', 'Meaning'],
    [
        ['Ticker',         'Click to jump to Decision Console for a full breakdown of that ticker.'],
        ['Regime',         'Market state, colour-coded (green / red / amber / grey).'],
        ['Action',         'Recommendation badge for that ticker.'],
        ['Conf %',         'Signal confidence score.'],
        ['Entry / Stop / Target / Shares', 'Suggested price levels for a trade.'],
        ['↻ (refresh)',    'Re-run the analysis for that one ticker only.'],
        ['📖 (book)',      'Save a trade plan (only visible for BUY or SELL SHORT signals).'],
    ],
    col_widths=[4, 11.5]
)

add_heading(doc, 'Summary chips', 2)
add_para(doc,
    'After analyzing, coloured chips appear above the table counting tickers by action state '
    '(e.g. BUY 3 · HOLD 5 · WAIT 4). This gives a quick market-wide overview at a glance.')

add_heading(doc, 'Column sorting', 2)
add_para(doc,
    'Click any column header to sort ascending / descending. Useful for surfacing the highest-'
    'confidence BUY signals or grouping all tickers in the same regime.')

# ══════════════════════════════════════════════════════════════════════════════
# 3. WATCHLIST HEAT MAP
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '3.  Watchlist Heat Map', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Route: /watchlist-heatmap', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'A visual card grid showing all watchlist tickers colour-coded by market regime and action '
    'signal. Designed for a quick at-a-glance market scan before diving into detail.')

add_heading(doc, 'How to use', 2)
add_bullet(doc, 'Select a date (defaults to today).')
add_bullet(doc, 'Click Scan Now — the engine scans all watchlist tickers and displays results as cards.')
add_bullet(doc, 'Click any card to open Decision Console for that ticker.')

add_heading(doc, 'Reading the cards', 2)
add_para(doc, 'Each card contains:')
add_bullet(doc, 'Ticker name — large text, coloured by regime.')
add_bullet(doc, 'Action badge — BUY, SHORT, SELL, COVER, HOLD, or WAIT.')
add_bullet(doc, 'Regime dot + label — market state with a coloured dot indicator.')
add_bullet(doc, 'Signal type — e.g. MA_CROSS, BREAKOUT, STOCH_CROSS.')
add_bullet(doc, 'Confidence bar — green ≥ 70%, amber 50–69%, red < 50%.')
add_bullet(doc, 'Price levels — Entry, Stop, Target (only when the signal is actionable).')
add_para(doc, 'Cards with an actionable signal (BUY or SELL SHORT) are highlighted with a coloured ring border.')

add_heading(doc, 'Filtering & sorting', 2)
add_table(doc,
    ['Control', 'Options'],
    [
        ['Filter bar', 'All  |  Actionable  |  BUY  |  SELL_SHORT  |  HOLD  |  WAIT'],
        ['Sort by',    'Confidence (default)  |  Ticker  |  Action  |  Regime'],
    ],
    col_widths=[4, 11.5]
)
add_note_box(doc, 'Use the Actionable filter to instantly narrow the grid to only tickers with entry signals.')

# ══════════════════════════════════════════════════════════════════════════════
# 4. DISCIPLINE DASHBOARD
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, '4.  Discipline Dashboard', 1, color=(0x2C, 0x5F, 0x4A))
add_para(doc, 'Route: /discipline', italic=True, color=(0x8F, 0x8C, 0x7F))
add_para(doc,
    'Monitors your trading discipline based on rules adapted from Jack Schwager\'s '
    'Market Wizards. It analyses your recent trade history and flags when you are '
    'breaking risk rules before a violation causes real damage.')

add_heading(doc, 'Stat cards (top row)', 2)
add_table(doc,
    ['Card', 'What it shows'],
    [
        ['Loss Streak',       'Number of consecutive recent losing trades.'],
        ['Streak Active',     'YES / No — whether you have hit your loss-streak threshold. When YES, position sizes are automatically reduced.'],
        ['Drawdown %',        'Current drawdown from your equity peak. A 20% drawdown triggers a mandatory-review rule.'],
        ['Rules Triggered',   'How many Schwager rules are currently violated out of the full rulebook.'],
    ],
    col_widths=[4, 11.5]
)
add_para(doc, 'Cards with a yellow border indicate a warning condition that needs attention.')

add_heading(doc, 'Last 10 Actions', 2)
add_para(doc, 'A sequence of your 10 most recent trade actions, colour-coded:')
add_bullet(doc, 'Green  = entry trade (BUY / SELL SHORT)')
add_bullet(doc, 'Red    = exit trade (SELL / COVER)')
add_bullet(doc, 'Grey   = neutral (HOLD / WAIT)')

add_heading(doc, 'Active Violations', 2)
add_table(doc,
    ['Severity', 'Meaning'],
    [
        ['🔴  BLOCK', 'A hard rule violation. The Decision Console will automatically reject BUY / SELL SHORT signals while this is active.'],
        ['🟡  WARN',  'A soft caution. Trading is still allowed but the engine will flag the risk in the reason chain.'],
    ],
    col_widths=[3, 12.5]
)

add_heading(doc, 'Schwager Rules panel', 2)
add_para(doc,
    'The full rulebook is shown with each rule either triggered (highlighted with colour) or '
    'clear (faded). For each rule you can see:')
add_bullet(doc, 'Rule ID and chapter reference from Market Wizards')
add_bullet(doc, 'The rule text')
add_bullet(doc, 'A detail message explaining why it was triggered (when active)')

add_heading(doc, 'When to check this screen', 2)
add_bullet(doc, 'Before running a scan — BLOCK violations will suppress entry signals automatically.')
add_bullet(doc, 'After a losing streak — the dashboard confirms whether position-size reduction is active.')
add_bullet(doc, 'After a large drawdown — the 20% drawdown rule triggers a mandatory strategy review.')

# ══════════════════════════════════════════════════════════════════════════════
# RECOMMENDED WORKFLOW
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Recommended Daily Workflow', 1)
steps = [
    ('Step 1', 'Discipline Dashboard',  'Check for any BLOCK violations. If present, resolve the underlying issue before trading.'),
    ('Step 2', 'Watchlist Heat Map',    'Run a scan for today\'s date. Use the Actionable filter to find entry signals.'),
    ('Step 3', 'Decision Console',      'Click into each promising card. Review regime, signal, reason chain, and R:R ratio.'),
    ('Step 4', 'Save Trade Plan',       'If the signal meets your criteria (confidence ≥ 60%, R:R ≥ 1.5:1), save a trade plan.'),
    ('Step 5', 'Decision Board',        'Check the full board at end-of-day for any regime changes across your watchlist.'),
]
add_table(doc,
    ['Step', 'Screen', 'Action'],
    steps,
    col_widths=[2, 4.5, 9]
)

# ══════════════════════════════════════════════════════════════════════════════
# TIPS
# ══════════════════════════════════════════════════════════════════════════════
add_heading(doc, 'Tips & Best Practices', 1)
tips = [
    ('Historical analysis',   'Use any past date in the date field to test how signals looked on historical bars.'),
    ('Confidence threshold',  'Signals below 60% are low-conviction. Consider waiting for a stronger setup.'),
    ('R:R ratio',             'Only take trades with R:R ≥ 1.5:1. The engine warns you below 1:1.'),
    ('HOLD vs WAIT',          'HOLD = you have an open position, keep it. WAIT = no position and no entry signal yet.'),
    ('Pattern confirmation',  'A detected chart pattern that agrees with the action signal (e.g. Double Bottom + BUY) increases conviction significantly.'),
    ('Position sizing',       'Set your Account Equity and Risk % accurately in Risk Preferences — the Shares figure depends on these values.'),
    ('Loss streak',           'If Streak Active shows YES, the engine already applies reduced position sizes automatically.'),
]
add_table(doc,
    ['Tip', 'Detail'],
    tips,
    col_widths=[4.5, 11]
)

# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════
out_path = '/home/tuanph/all_projects/stockpredict/Decision_Helper_User_Guide.docx'
doc.save(out_path)
print(f'Saved: {out_path}')
