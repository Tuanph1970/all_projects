"""Generate road_map_version_11.xlsx — plan for the Trader Decision Helper feature.

Inputs synthesized into this workbook:
  - /home/tuanph/TuanPH/all_projects/knowledge/stock_pipeline_v3.drawio  (pipeline reference)
  - /home/tuanph/TuanPH/all_projects/knowledge/stock/books/getting-started-in-technical-analysis_compress.pdf/
    parts 1-4, chapter notes / index / README files (Schwager: "Getting Started in Technical Analysis")
  - Existing stockpredict codebase inventory (services, frontend pages, indicators, signal logic)

Run: `python3 build_road_map_v11.py`  -> writes road_map_version_11.xlsx in the same folder.
This script is idempotent; re-run after editing rows below to regenerate.
"""
from __future__ import annotations
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "road_map_version_11.xlsx")

# ---- Style helpers ---------------------------------------------------------
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1F4E78")
SUBTITLE_FONT = Font(name="Calibri", italic=True, size=10, color="595959")
H2_FONT = Font(name="Calibri", bold=True, size=12, color="1F4E78")

STATUS_FILLS = {
    "Done":         PatternFill("solid", fgColor="C6EFCE"),
    "In Progress":  PatternFill("solid", fgColor="FFEB9C"),
    "Partial":      PatternFill("solid", fgColor="FFEB9C"),
    "Stub":         PatternFill("solid", fgColor="FFC7CE"),
    "Planned":      PatternFill("solid", fgColor="DDEBF7"),
    "Not Started":  PatternFill("solid", fgColor="F2F2F2"),
    "Blocked":      PatternFill("solid", fgColor="FFC7CE"),
    "Proposed":     PatternFill("solid", fgColor="DDEBF7"),
    "Open":         PatternFill("solid", fgColor="FFEB9C"),
    "Answered":     PatternFill("solid", fgColor="C6EFCE"),
}
PRIORITY_FILLS = {
    "P0 (must)":   PatternFill("solid", fgColor="FFC7CE"),
    "P1 (should)": PatternFill("solid", fgColor="FFEB9C"),
    "P2 (could)":  PatternFill("solid", fgColor="DDEBF7"),
}


def style_header_row(ws, row, last_col):
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_table(ws, headers, rows, header_row=1, freeze=True):
    for i, h in enumerate(headers, 1):
        ws.cell(row=header_row, column=i, value=h)
    style_header_row(ws, header_row, len(headers))
    for r, row in enumerate(rows, start=header_row + 1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    if freeze:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_conditional(ws, col_letter, first_row, last_row, fills_map):
    rng = f"{col_letter}{first_row}:{col_letter}{last_row}"
    for value, fill in fills_map.items():
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill)
        )


# ===========================================================================
wb = Workbook()

# ---------------------------------------------------------------------------
# Sheet 1: Overview
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "Overview"

ws["A1"] = "StockPredict — Road Map v11: Trader Decision Helper"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A1:F1")

ws["A2"] = ("Plan for a new feature that helps traders decide WHEN to take an opportunity, WHEN to cut a loss, "
            "and HOW to size the position — operationalising Schwager's 'Getting Started in Technical Analysis'.")
ws["A2"].font = SUBTITLE_FONT
ws.merge_cells("A2:F2")

overview_rows = [
    ("Goal",
     "Give the trader, for any watched ticker on any bar, a single recommendation (BUY / SELL / "
     "HOLD / WAIT / REVERSE) plus a stop, a target, and a position size — all justified by an "
     "auditable chain of indicators, patterns, regime, and confirmation rules drawn from Schwager."),
    ("Non-goal (v1)",
     "We do NOT auto-execute orders. The helper recommends; the trader (or the broker integration "
     "in a later phase) acts. Paper-trading + journal-loop closes the feedback cycle."),
    ("Inputs",
     "1) Pipeline diagram: knowledge/stock_pipeline_v3.drawio  "
     "2) Book notes: knowledge/stock/books/getting-started-in-technical-analysis_compress.pdf/part_1..4/  "
     "3) Existing stockpredict services (9 FastAPI services + React 19 frontend)."),
    ("New service",
     "decision-helper-service (port 8010). Orchestrates: regime -> signal -> confirm -> "
     "failed-signal check -> state-machine update -> recommendation + stop/target/size. Emits "
     "events to journal-service and portfolio-service."),
    ("Existing services to extend",
     "analysis-service (add ATR, swing pivots, regime endpoint, midtrend-entry detectors); "
     "ai-insights-service (add H&S / double-top / flag-pennant / WRD / gaps / trap detectors); "
     "journal-service (auto-pre-fill draft entries from helper events); "
     "portfolio-service (sync stops + targets as alerts)."),
    ("New frontend pages",
     "DecisionConsole (per-symbol), WatchlistDecisionBoard (multi-symbol), TradePlanModal "
     "(gated entry — refuses to save without stop+target), BacktestStudio, DisciplineDashboard."),
    ("How to read this workbook",
     "Pipeline_Mapping → see every drawio stage mapped to existing code and new modules.  "
     "Architecture → components to build.  Backlog → concrete tickets.  Phased_Roadmap → "
     "order of attack.  Book_Chapter_Map → which chapter informs which module.  "
     "Risks_Open_Questions → what we need to clarify before coding."),
    ("Success criteria",
     "1) For 5 tickers on daily bars, helper produces a recommendation in < 2s.  "
     "2) Walk-forward backtest on 3+ non-overlapping periods shows positive expectancy "
     "after slippage.  3) Every entry recorded in journal includes auto-generated reason "
     "linked to the signal+regime that produced it.  4) Trader cannot save a trade plan "
     "without a stop and a target (Ch 16-17 discipline)."),
    ("Out of scope (v1)",
     "Intraday/tick data, order-execution integration with brokers, options/futures, "
     "automated portfolio rebalancing, multi-account."),
]
ws["A4"] = "What this is"
ws["A4"].font = H2_FONT
for i, (k, v) in enumerate(overview_rows, start=5):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True)
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=i, column=2, value=v).alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)

ws.column_dimensions["A"].width = 26
for col in "BCDEF":
    ws.column_dimensions[col].width = 24

# ---------------------------------------------------------------------------
# Sheet 2: Pipeline_Mapping  — every stage in stock_pipeline_v3.drawio
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Pipeline_Mapping")
headers = ["Stage ID", "Stage (drawio)", "What it does", "Current state in stockpredict",
           "New module to build", "Owner service", "Book chapter(s)", "Priority", "Status"]

pipeline_rows = [
    ("s1_universe", "1. Select 5 stock tickers",
     "Pre-filter universe (fundamentals, liquidity) before technical analysis runs.",
     "Watchlist + screener exist (portfolio-service, analysis-service)",
     "Wire watchlist -> decision pipeline trigger; add liquidity filter helper.",
     "portfolio-service + decision-helper-service",
     "Ch 16 (market selection)", "P1 (should)", "Planned"),

    ("s2_data", "2. Get historical OHLCV (daily)",
     "Daily OHLCV per ticker; splits/dividend adjusted.",
     "Exists in stock-data-service (vnstock)",
     "Expose a clean daily-bars endpoint that includes adjusted close + volume.",
     "stock-data-service",
     "Ch 13 (data)", "P0 (must)", "Partial"),

    ("s3_indicators", "3. Compute indicators",
     "MAs, ATR, RSI/MACD/Stoch, swing pivots — every bar.",
     "RSI, MACD, BB, SMA/EMA in analysis-service; ADX in prediction-service (ta-lib)",
     "Add ATR endpoint, swing-pivot detector, consolidated indicator pack response.",
     "analysis-service",
     "Ch 3, 6", "P0 (must)", "Partial"),

    ("s4_patterns", "4. Detect chart patterns",
     "Gaps, continuation, reversal, wide-range bar.",
     "ai-insights-service: Uptrend/Downtrend/Consolidation/Breakout only",
     "Extend pattern catalog: H&S, double top/bottom, flag, pennant, triangle, WRD, gaps, "
     "bull/bear traps.",
     "ai-insights-service",
     "Ch 5, 11, 12", "P0 (must)", "Partial"),

    ("s5_regime", "5. Regime? (trend vs trading range)",
     "Classify the bar as trending or ranging using MA slope, swing structure, ADX, range/ATR.",
     "Nothing explicit — must build",
     "regime_classifier module — outputs {trend, range, undefined} + confidence.",
     "decision-helper-service",
     "Ch 3, 4", "P0 (must)", "Not Started"),

    ("s6L_direction", "6L. Up or down trend?",
     "Branch within trend regime: direction = long-bias or short-bias.",
     "Inferred from SMA/EMA cross today",
     "trend_direction helper using long-MA slope + swing structure.",
     "decision-helper-service",
     "Ch 3", "P0 (must)", "Not Started"),

    ("s6R_range", "6R. Trading range — locate S/R",
     "Inside a range, find support floor and resistance ceiling.",
     "analysis-service _find_support_resistance() (local minmax)",
     "Promote S/R into a typed level object with strength/age; expose to decision helper.",
     "analysis-service",
     "Ch 4", "P0 (must)", "Partial"),

    ("s7L_signal", "8L. Trend-following signal",
     "MA crossover or N-day breakout entry trigger.",
     "Score-based signal in analysis-service _determine_signal() — not state-aware",
     "signal_engine.trend: MA-cross, N-day breakout, midtrend pullback/retracement/flag.",
     "decision-helper-service",
     "Ch 8, 14", "P0 (must)", "Not Started"),

    ("s7R_breakout", "7R. Breakout?",
     "Has the range broken out (price closed past S or R)?",
     "Nothing explicit",
     "breakout_detector with penetration + close-confirmation.",
     "decision-helper-service",
     "Ch 4, 14", "P0 (must)", "Not Started"),

    ("s9L_filter", "9L. Filter + confirmation (trend)",
     "Direction filter, penetration, time delay before entering.",
     "Nothing explicit",
     "confirmation_layer: penetration N%, time-delay K bars, multi-TF alignment.",
     "decision-helper-service",
     "Ch 7, 14", "P0 (must)", "Not Started"),

    ("s8R_confirm", "8R. Confirm breakout direction",
     "Penetration + time-delay confirmation for breakout entry.",
     "Nothing explicit",
     "confirmation_layer reused on breakout branch.",
     "decision-helper-service",
     "Ch 7, 14", "P0 (must)", "Not Started"),

    ("s10L_failed", "10L. Signal failed? -> consider reversal",
     "Detect a failed trend signal; if yes, reverse (Schwager: most important rule).",
     "Nothing exists",
     "failed_signal_detector for trend branch: bull/bear trap, false trend-line break, "
     "filled gap, return to spike/WRD extreme.",
     "decision-helper-service",
     "Ch 11", "P0 (must)", "Not Started"),

    ("s9R_failed", "9R. Breakout failed?",
     "Bull / bear trap detection on the range branch.",
     "Nothing exists",
     "failed_signal_detector reused on range branch.",
     "decision-helper-service",
     "Ch 11", "P0 (must)", "Not Started"),

    ("state_machine", "Position state machine (FLAT / LONG / SHORT)",
     "Per (user, ticker): current state + incoming signal -> next action. Pyramid self-loops "
     "(Ch 8) are deferred.",
     "Portfolio tracks holdings only — no state-machine for active 'trade idea'",
     "position_state_machine — DB-backed; transitions on confirmed signal, stop, target, "
     "or reversal.",
     "decision-helper-service",
     "Ch 8, 14", "P0 (must)", "Not Started"),

    ("out_actions", "BUY / SELL / SELL short / COVER / HOLD / WAIT",
     "Action emitted at each bar based on (state, signal).",
     "Nothing exists",
     "decision_orchestrator — outputs the action plus reason chain.",
     "decision-helper-service",
     "Ch 12, 14, 17", "P0 (must)", "Not Started"),

    ("log_loop", "Log decision (action + reason + state) & advance bar",
     "Persist what was decided and why; backtest replay-friendly.",
     "Journal logs trades; no per-bar decision log",
     "decision_runs table + Kafka/RabbitMQ event 'decision.recommendation.created'.",
     "decision-helper-service",
     "Ch 16", "P0 (must)", "Not Started"),

    ("risk_stops", "Risk overlay: Stop-loss (Ch 9)",
     "Initial + trailing stops anchored on ATR / swing / S/R / pattern boundary.",
     "Portfolio price alerts only — no auto stop computation",
     "stop_module: chart-based, swing, ATR, pattern, WRD, money-stop, trailing. "
     "Validation: never widen.",
     "decision-helper-service + portfolio-service (sync as alert)",
     "Ch 9, 17", "P0 (must)", "Not Started"),

    ("risk_sizing", "Risk overlay: Position sizing (Ch 16)",
     "Compute size from (account risk %, entry - stop). ATR-adjusted; equity-adjusted.",
     "PositionCalculator.js UI exists but is not wired",
     "sizing_module: fixed-fractional, vol-adjusted, equity-adjusted, correlation cap, "
     "losing-streak adjustment.",
     "decision-helper-service",
     "Ch 16", "P0 (must)", "Not Started"),

    ("risk_targets", "Risk overlay: Profit targets (Ch 10)",
     "S/R level, measured-move, partial exits + trailing take-profit.",
     "portfolio-service has target_price field on alerts; nothing computes it",
     "target_module: S/R objective, measured move, RSI extreme (range), trailing TP.",
     "decision-helper-service + portfolio-service (sync)",
     "Ch 10", "P0 (must)", "Not Started"),

    ("regime_inputs", "Regime classifier inputs",
     "Long-MA slope, swing high/low structure, range width vs ATR, ADX, pattern context.",
     "ADX exists; swing structure does not",
     "Aggregate features into a single 'regime feature vector'.",
     "analysis-service -> decision-helper-service",
     "Ch 3, 4", "P0 (must)", "Not Started"),

    ("osc_note", "Oscillator-use rule (Ch 6)",
     "In trend -> oscillators are alert-only. In range -> actionable at boundaries.",
     "Score-based signal blends oscillators with price unconditionally",
     "confirmation_layer applies this rule by branch (trend vs range).",
     "decision-helper-service",
     "Ch 6", "P1 (should)", "Not Started"),

    ("failed_signal_principle", "Failed-signal principle (Ch 11)",
     "Schwager's 'most important rule' — a failed pattern is itself a stronger signal.",
     "Nothing exists",
     "failed_signal_detector wires into BOTH branches and triggers reversal.",
     "decision-helper-service",
     "Ch 11, 17 (rule 76)", "P0 (must)", "Not Started"),

    ("backtest_harness", "Backtest / walk-forward (Ch 15)",
     "Replay the pipeline on history with realistic slippage; out-of-sample evaluation.",
     "prediction-service has simulation/backtest UI; no walk-forward for decision helper",
     "backtest_harness reusing decision_orchestrator over historical bars; multi-period, "
     "multi-market.",
     "decision-helper-service + prediction-service",
     "Ch 15", "P1 (should)", "Not Started"),

    ("discipline", "Discipline overlay (Ch 17, 82 rules)",
     "Reject trade plans without stop+target; detect averaging-down, stop-widening, "
     "post-streak risk creep.",
     "Journal captures emotional state; no enforcement layer",
     "discipline_checks module + TradePlanModal gate + DisciplineDashboard.",
     "decision-helper-service + journal-service + frontend",
     "Ch 16, 17", "P1 (should)", "Not Started"),
]
write_table(ws, headers, pipeline_rows)
autosize(ws, [16, 28, 38, 38, 42, 30, 16, 14, 14])
add_conditional(ws, "I", 2, 1 + len(pipeline_rows), STATUS_FILLS)
add_conditional(ws, "H", 2, 1 + len(pipeline_rows), PRIORITY_FILLS)

# ---------------------------------------------------------------------------
# Sheet 3: Architecture
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Architecture")
headers = ["Component", "Type", "Location (proposed)", "Purpose",
           "Depends on", "Book chapters", "Status"]

arch_rows = [
    # New service
    ("decision-helper-service", "Service (new)", "stock_microservices/services/decision-helper-service/",
     "Orchestrates the full pipeline; emits recommendation per (user, ticker, bar). FastAPI on port 8010.",
     "stock-data-service, analysis-service, ai-insights-service, journal-service, portfolio-service, MySQL, Redis, RabbitMQ",
     "All", "Not Started"),
    ("regime_classifier", "Module", ".../app/services/regime_classifier.py",
     "Classifies bar as TREND / RANGE / UNDEFINED using MA slope, swing structure, ADX, range/ATR.",
     "indicators pack, swing pivots", "Ch 3, 4", "Not Started"),
    ("signal_engine", "Module", ".../app/services/signal_engine/",
     "Submodules: trend (MA cross, N-day breakout, midtrend), range (S/R bounce, breakout), countertrend (osc fade).",
     "regime_classifier, indicators", "Ch 6, 8, 14", "Not Started"),
    ("confirmation_layer", "Module", ".../app/services/confirmation.py",
     "Penetration %, time delay K bars, multi-TF alignment, oscillator-as-alert rule.",
     "signal_engine output", "Ch 7, 14", "Not Started"),
    ("failed_signal_detector", "Module", ".../app/services/failed_signal.py",
     "12 detectors from Ch 11 (bull/bear trap, false TL break, filled gap, etc.). Triggers reversal.",
     "patterns, signals, state_machine", "Ch 11", "Not Started"),
    ("stop_module", "Module", ".../app/services/stop_module.py",
     "Chart-based, swing, ATR, pattern, WRD, money-stop, trailing. Validation: stops never widen.",
     "indicators, patterns, state_machine", "Ch 9", "Not Started"),
    ("target_module", "Module", ".../app/services/target_module.py",
     "S/R objective, measured move, RSI extreme (range), trailing TP. Partial exit support.",
     "patterns, S/R levels", "Ch 10", "Not Started"),
    ("sizing_module", "Module", ".../app/services/sizing_module.py",
     "Fixed-fractional, ATR-adjusted, equity-adjusted, correlation cap, losing-streak adjustment.",
     "stop_module, account state, position correlations", "Ch 16", "Not Started"),
    ("position_state_machine", "Module", ".../app/services/state_machine.py",
     "Per (user, ticker): FLAT / LONG / SHORT. Persists transitions; emits events.",
     "decision_orchestrator", "Ch 8, 14", "Not Started"),
    ("decision_orchestrator", "Module", ".../app/services/orchestrator.py",
     "Runs the pipeline for one bar of one ticker. Returns recommendation + reason chain.",
     "everything above", "Ch 12, 14", "Not Started"),
    ("discipline_checks", "Module", ".../app/services/discipline.py",
     "Plan-before-trade enforcement; no-averaging-down; stop-widening detector; streak detectors.",
     "journal data, account history", "Ch 16, 17", "Not Started"),
    ("backtest_harness", "Module", ".../app/services/backtest.py",
     "Walk-forward replay of orchestrator over historical bars; slippage; out-of-sample report.",
     "decision_orchestrator, historical OHLCV", "Ch 15", "Not Started"),
    # DB tables
    ("decision_runs (table)", "DB table", "MySQL `stockpredict`.decision_runs",
     "One row per (user, ticker, bar): regime, signal, confirmation, recommendation, "
     "reason_json, stop, target, size, state_before, state_after.",
     "MySQL", "All", "Not Started"),
    ("position_states (table)", "DB table", "MySQL `stockpredict`.position_states",
     "Current state per (user, ticker): state, entry_bar, entry_price, current_stop, current_target, size.",
     "MySQL", "Ch 8, 14", "Not Started"),
    ("signal_history (table)", "DB table", "MySQL `stockpredict`.signal_history",
     "Append-only log of raw signals (pre-confirmation) for analysis and backtest replay.",
     "MySQL", "Ch 15", "Not Started"),
    ("user_risk_prefs (table)", "DB table", "MySQL `stockpredict`.user_risk_prefs",
     "Per-user: risk_per_trade_pct, max_concurrent_positions, max_correlation, account_equity_snapshot.",
     "MySQL, auth-service", "Ch 16", "Not Started"),
    # Events
    ("decision.recommendation.created", "Event", "RabbitMQ /stock",
     "Published every time a new recommendation is emitted; consumed by frontend (WebSocket) "
     "and journal-service (draft entry).",
     "RabbitMQ", "Ch 16", "Not Started"),
    ("decision.signal.failed", "Event", "RabbitMQ /stock",
     "Published when a signal turns into a failed-signal reversal candidate.",
     "RabbitMQ", "Ch 11", "Not Started"),
    ("decision.stop.triggered / decision.target.hit", "Event", "RabbitMQ /stock",
     "Published when a tracked position hits its stop or target. Consumed by journal-service "
     "and portfolio-service.",
     "RabbitMQ", "Ch 9, 10", "Not Started"),
    # Existing service extensions
    ("analysis-service: ATR + swings", "Extension", "stock_microservices/services/analysis-service/app/services/analysis_service.py",
     "Add _calculate_atr(), _find_swing_pivots(), and a /analysis/regime/{symbol} endpoint.",
     "OHLCV", "Ch 3, 9", "Not Started"),
    ("ai-insights-service: pattern v2", "Extension", "stock_microservices/services/ai-insights-service/app/services/insights_service.py",
     "Add H&S, double top/bottom, flag, pennant, triangle, gap, WRD, bull/bear trap detectors.",
     "OHLCV", "Ch 5, 11, 12", "Not Started"),
    ("journal-service: auto-draft", "Extension", "stock_microservices/services/journal-service/app/services/diary_service.py",
     "Subscribe to decision events; create draft journal entries pre-filled with reason / "
     "stop / target / size. Trader confirms or edits.",
     "RabbitMQ, decision-helper-service", "Ch 16", "Not Started"),
    ("portfolio-service: sync alerts", "Extension", "stock_microservices/services/portfolio-service/app/services/portfolio_service.py",
     "When recommendation produces a stop/target, materialise as a portfolio alert "
     "automatically; cancel when state transitions.",
     "RabbitMQ", "Ch 9, 10", "Not Started"),
    # Frontend
    ("DecisionConsole", "Frontend page", "stock_react_javascript/stock-react-app/src/components/Pages/DecisionConsole/",
     "Per-symbol: chart + regime ribbon + signal card + recommendation + stop/target/size + reason chain.",
     "decision-helper-service API", "Ch 12", "Not Started"),
    ("WatchlistDecisionBoard", "Frontend page", ".../components/Pages/WatchlistDecisionBoard/",
     "Multi-symbol table: regime, signal, recommendation, change since last bar. Sortable/filterable.",
     "decision-helper-service API", "Ch 16", "Not Started"),
    ("TradePlanModal", "Frontend component", ".../components/Common/TradePlanModal.js",
     "Gated entry: cannot save a planned trade without stop + target. Pre-fills from recommendation.",
     "journal-service, decision-helper-service", "Ch 9, 16, 17", "Not Started"),
    ("BacktestStudio", "Frontend page", ".../components/Pages/BacktestStudio/",
     "Walk-forward UI: pick periods + markets + parameter sets, view OOS expectancy, drawdown, "
     "win rate.",
     "decision-helper-service /backtest API", "Ch 15", "Not Started"),
    ("DisciplineDashboard", "Frontend page", ".../components/Pages/DisciplineDashboard/",
     "Streaks, drawdown, rule-violation log, post-loss break suggestion.",
     "discipline_checks API, journal-service", "Ch 16, 17", "Not Started"),
]
write_table(ws, headers, arch_rows)
autosize(ws, [38, 16, 56, 60, 50, 16, 14])
add_conditional(ws, "G", 2, 1 + len(arch_rows), STATUS_FILLS)

# ---------------------------------------------------------------------------
# Sheet 4: Feature_Backlog
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Feature_Backlog")
headers = ["ID", "Phase", "Area", "Feature", "Description / Acceptance criteria",
           "Chapter(s)", "Effort", "Priority", "Depends on", "Status"]

backlog = [
    # Phase A: Foundations
    ("F-001", "A", "analysis-service", "Add ATR indicator",
     "Implement _calculate_atr(period=14). Expose at /analysis/atr/{symbol}. "
     "AC: returns 14-bar ATR aligned with OHLCV index; unit test on 2 known series.",
     "Ch 9", "S", "P0 (must)", "—", "Not Started"),
    ("F-002", "A", "analysis-service", "Add swing-pivot detector",
     "Detect higher highs / lower lows over rolling window K. Expose at /analysis/swings/{symbol}. "
     "AC: returns list of (bar_index, type, price); deterministic across reruns.",
     "Ch 3", "M", "P0 (must)", "—", "Not Started"),
    ("F-003", "A", "analysis-service", "Promote S/R levels to typed objects",
     "Refactor _find_support_resistance() to return {price, strength, age_bars, touches}. "
     "AC: backwards-compatible response; new fields populated.",
     "Ch 4", "S", "P0 (must)", "F-002", "Not Started"),
    ("F-004", "A", "analysis-service", "Regime feature vector endpoint",
     "/analysis/regime/{symbol} -> {regime, confidence, features:{slope, hh_hl_count, adx, range_atr_ratio}}.",
     "Ch 3, 4", "M", "P0 (must)", "F-001, F-002", "Not Started"),
    ("F-005", "A", "ai-insights-service", "Extend pattern catalog",
     "Add detectors: H&S, double top/bottom, flag, pennant, triangle, gap, wide-ranging day. "
     "AC: each returns {found, confidence, key_bars, measured_move_target}.",
     "Ch 5, 12", "L", "P0 (must)", "F-002", "Not Started"),
    ("F-006", "A", "ai-insights-service", "Bull / bear trap detector",
     "Detect penetration of S/R that closes back inside range within N bars. AC: returns "
     "trap event with severity.",
     "Ch 11", "M", "P0 (must)", "F-003, F-005", "Not Started"),

    # Phase B: Decision helper service skeleton + signal engine
    ("F-010", "B", "decision-helper-service", "Bootstrap FastAPI service (port 8010)",
     "Stand up service with health probe, DB migrations baseline, Traefik route /helper/*. "
     "AC: deploy.sh starts it; GET /health 200.",
     "—", "S", "P0 (must)", "—", "Not Started"),
    ("F-011", "B", "decision-helper-service", "DB schema: decision_runs, position_states, signal_history, user_risk_prefs",
     "Alembic migration with the four tables described in Architecture sheet.",
     "—", "M", "P0 (must)", "F-010", "Not Started"),
    ("F-012", "B", "decision-helper-service", "regime_classifier module",
     "Pull regime features from analysis-service; output {trend, range, undefined} + confidence. "
     "AC: unit test on synthetic trending vs ranging series.",
     "Ch 3, 4", "M", "P0 (must)", "F-004, F-010", "Not Started"),
    ("F-013", "B", "decision-helper-service", "signal_engine.trend.ma_cross",
     "MA-crossover detector with configurable fast/slow periods. AC: emits BUY/SELL signal "
     "at crossover bar.",
     "Ch 14", "S", "P0 (must)", "F-010", "Not Started"),
    ("F-014", "B", "decision-helper-service", "signal_engine.trend.n_day_breakout",
     "N-day high/low breakout detector. AC: configurable N; long signal on close > N-day high.",
     "Ch 14", "S", "P0 (must)", "F-010", "Not Started"),
    ("F-015", "B", "decision-helper-service", "signal_engine.range.sr_bounce_breakout",
     "Inside a range: bounce signal at support/resistance; breakout signal on close past S/R.",
     "Ch 4, 14", "M", "P0 (must)", "F-003", "Not Started"),
    ("F-016", "B", "decision-helper-service", "confirmation_layer",
     "Penetration % + time-delay K + multi-TF alignment + oscillator-as-alert rule. "
     "Configurable thresholds. AC: confirms/rejects each raw signal with reason.",
     "Ch 7, 14", "M", "P0 (must)", "F-013, F-014, F-015", "Not Started"),

    # Phase C: Risk overlay + state machine
    ("F-020", "C", "decision-helper-service", "stop_module — chart / swing / ATR / pattern",
     "Compute initial stop per entry type. Validation: stop never widens after entry "
     "(unit test the CIC anti-pattern).",
     "Ch 9", "L", "P0 (must)", "F-001, F-002, F-003", "Not Started"),
    ("F-021", "C", "decision-helper-service", "stop_module — trailing stop",
     "Update stop on each new favorable bar; never relax.",
     "Ch 9, 10", "M", "P0 (must)", "F-020", "Not Started"),
    ("F-022", "C", "decision-helper-service", "target_module — S/R + measured move",
     "Pick first viable target from S/R level or measured move from pattern.",
     "Ch 10", "M", "P0 (must)", "F-003, F-005", "Not Started"),
    ("F-023", "C", "decision-helper-service", "target_module — partial exits",
     "Support 50% at first target + trailing on remainder.",
     "Ch 10", "M", "P1 (should)", "F-022", "Not Started"),
    ("F-024", "C", "decision-helper-service", "sizing_module v1 (fixed-fractional)",
     "size = floor(risk_pct * equity / (entry - stop)). AC: rejects size < 1.",
     "Ch 16", "S", "P0 (must)", "F-020", "Not Started"),
    ("F-025", "C", "decision-helper-service", "sizing_module v2 (ATR + equity + correlation + streak)",
     "ATR-adjustment, equity-adjustment, correlation cap, losing-streak halving.",
     "Ch 16", "L", "P1 (should)", "F-024", "Not Started"),
    ("F-026", "C", "decision-helper-service", "position_state_machine",
     "FLAT/LONG/SHORT transitions; persists to position_states. "
     "AC: cannot transition LONG->SHORT without passing through FLAT (v1).",
     "Ch 8, 14", "M", "P0 (must)", "F-011", "Not Started"),
    ("F-027", "C", "decision-helper-service", "decision_orchestrator",
     "Run one bar of one ticker through the pipeline; return recommendation + reason chain; "
     "persist to decision_runs.",
     "Ch 12, 14", "L", "P0 (must)", "F-012..F-026", "Not Started"),

    # Phase D: Frontend
    ("F-030", "D", "frontend", "API client + decision context",
     "Add helperApi.js; React context for current symbol's latest recommendation.",
     "—", "S", "P0 (must)", "F-027", "Not Started"),
    ("F-031", "D", "frontend", "DecisionConsole page",
     "Per-symbol: chart, regime ribbon, signal card, recommendation, stop/target/size, "
     "reason chain. Manual test on 3 tickers in browser.",
     "Ch 12", "L", "P0 (must)", "F-030", "Not Started"),
    ("F-032", "D", "frontend", "WatchlistDecisionBoard",
     "Multi-symbol table of recommendations; sort by signal-change-since-last-bar.",
     "Ch 16", "M", "P0 (must)", "F-031", "Not Started"),
    ("F-033", "D", "frontend", "TradePlanModal (gated)",
     "Cannot save plan without stop + target; pre-fills from recommendation. "
     "Wires to journal-service.",
     "Ch 9, 16, 17", "M", "P0 (must)", "F-031", "Not Started"),
    ("F-034", "D", "journal-service", "Auto-draft journal entry from decision event",
     "Subscribe to decision.recommendation.created; create draft entry with reason chain.",
     "Ch 16", "M", "P1 (should)", "F-027", "Not Started"),
    ("F-035", "D", "portfolio-service", "Sync stops + targets as portfolio alerts",
     "When recommendation has a stop/target, create matching alert. Cancel on state exit.",
     "Ch 9, 10", "M", "P1 (should)", "F-027", "Not Started"),

    # Phase E: Failed signals + midtrend + pyramiding
    ("F-040", "E", "decision-helper-service", "failed_signal_detector — 12 detectors",
     "Implement each Ch-11 type as boolean + parameters (bull/bear trap, false TL break, "
     "filled gap, return to spike, WRD return, counter-anticipated flag, opposite breakout, "
     "H&S target failure, double-top break, rounding break, future-reliability tag).",
     "Ch 11", "L", "P0 (must)", "F-005, F-006", "Not Started"),
    ("F-041", "E", "decision-helper-service", "Reversal action wiring",
     "When detector fires, orchestrator emits REVERSE; state-machine flips LONG<->SHORT via FLAT.",
     "Ch 11, 17", "M", "P0 (must)", "F-040, F-026", "Not Started"),
    ("F-042", "E", "decision-helper-service", "signal_engine.midtrend",
     "Four midtrend entries: % retracement, minor-reaction reversal, continuation breakout, "
     "long-MA reaction.",
     "Ch 8", "L", "P1 (should)", "F-027", "Not Started"),
    ("F-043", "E", "decision-helper-service", "Pyramiding (deferred guard rails)",
     "Allow additive entries to a winning position with 3 risk guards: max 50% from pyramids, "
     "tighter stops on add, no-add after N consecutive losses.",
     "Ch 8, 14", "L", "P2 (could)", "F-042, F-025", "Not Started"),

    # Phase F: Backtest harness
    ("F-050", "F", "decision-helper-service", "backtest_harness — replay engine",
     "Run orchestrator over historical bars; produce trade ledger + equity curve.",
     "Ch 15", "L", "P1 (should)", "F-027", "Not Started"),
    ("F-051", "F", "decision-helper-service", "Walk-forward train/test split",
     "Non-overlapping windows; optimise on train, evaluate on test; aggregate across windows.",
     "Ch 15", "L", "P1 (should)", "F-050", "Not Started"),
    ("F-052", "F", "decision-helper-service", "Slippage + commissions model",
     "Configurable bps slippage and per-trade commission; applied in backtest.",
     "Ch 15", "S", "P1 (should)", "F-050", "Not Started"),
    ("F-053", "F", "decision-helper-service", "Parameter-set sensitivity report",
     "Run 5+ neighboring parameter sets; report variance — flag fragile optima.",
     "Ch 15", "M", "P2 (could)", "F-051", "Not Started"),
    ("F-054", "F", "frontend", "BacktestStudio page",
     "UI for parameter sets / periods / markets; render equity curve + drawdown + win rate.",
     "Ch 15", "L", "P2 (could)", "F-051", "Not Started"),

    # Phase G: Discipline
    ("F-060", "G", "decision-helper-service", "discipline_checks — plan-before-trade",
     "Block any TradePlanModal save where stop or target is missing.",
     "Ch 16, 17", "S", "P1 (should)", "F-033", "Not Started"),
    ("F-061", "G", "decision-helper-service", "discipline_checks — no averaging down",
     "Detect when an open losing position receives a new entry signal in the same direction; "
     "warn / require override.",
     "Ch 17 (rule 14)", "M", "P1 (should)", "F-026", "Not Started"),
    ("F-062", "G", "decision-helper-service", "discipline_checks — stop-widening detector",
     "Warn whenever user-edited stop is wider than current orchestrator stop.",
     "Ch 9, 17", "S", "P1 (should)", "F-020", "Not Started"),
    ("F-063", "G", "decision-helper-service", "Streak detectors + post-loss break",
     "Detect N consecutive losses or > 20% drawdown; trigger size reduction; suggest break.",
     "Ch 16, 17", "M", "P2 (could)", "F-025", "Not Started"),
    ("F-064", "G", "frontend", "DisciplineDashboard",
     "Streaks, drawdown, rule-violation log, last-N-trades summary.",
     "Ch 16, 17", "M", "P2 (could)", "F-060..F-063", "Not Started"),
    ("F-065", "G", "all", "Schwager 82-rules engine (advisory)",
     "Codify rules from Ch 17 as advisory checks surfaced in DecisionConsole sidebar.",
     "Ch 17", "L", "P2 (could)", "F-031", "Not Started"),
]
write_table(ws, headers, backlog)
autosize(ws, [10, 8, 22, 36, 70, 16, 9, 14, 22, 14])
add_conditional(ws, "J", 2, 1 + len(backlog), STATUS_FILLS)
add_conditional(ws, "H", 2, 1 + len(backlog), PRIORITY_FILLS)

# ---------------------------------------------------------------------------
# Sheet 5: Phased_Roadmap
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Phased_Roadmap")
headers = ["Phase", "Name", "Goal", "Key deliverables", "Exit criteria",
           "Est. weeks", "Depends on", "Status"]

phases = [
    ("A", "Foundations: indicators + patterns + regime features",
     "Make the inputs the decision pipeline needs available from existing services.",
     "ATR + swing pivots (analysis-service); S/R typed levels; regime feature endpoint; "
     "extended pattern catalog + bull/bear trap detector (ai-insights-service).",
     "All Phase-A backlog (F-001..F-006) green. Decision service can fetch every feature it "
     "needs from existing services with one HTTP call per type.",
     "2", "—", "Not Started"),
    ("B", "decision-helper-service skeleton + signal engine v1",
     "Stand up the new service; ship trend-following + range signal generators with a "
     "confirmation layer.",
     "Service bootstrap, DB schema, regime_classifier, signal_engine (MA cross, N-day breakout, "
     "S/R bounce / breakout), confirmation_layer.",
     "F-010..F-016 green. End-to-end: given (user, ticker, bar) the service emits a confirmed "
     "raw signal with a reason payload.",
     "3", "A", "Not Started"),
    ("C", "Risk overlay + state machine + orchestrator",
     "Wrap signals with stop / target / size and a persistent FLAT/LONG/SHORT state machine.",
     "stop_module, target_module, sizing_module v1, position_state_machine, decision_orchestrator.",
     "F-020..F-027 green. Orchestrator returns a full recommendation (action + stop + target + size + reason) "
     "stored in decision_runs. Replay on 5 tickers, 1 year of bars completes < 30s.",
     "3", "B", "Not Started"),
    ("D", "Frontend MVP: DecisionConsole + Watchlist + TradePlan + journal/portfolio wiring",
     "Make the helper usable by a real trader; close the loop with journal and portfolio alerts.",
     "DecisionConsole, WatchlistDecisionBoard, TradePlanModal (gated), journal auto-draft, "
     "portfolio alert sync.",
     "Trader can review 5 watched tickers, accept a recommendation, save a plan (with stop+target "
     "enforced), and see the corresponding journal draft + portfolio alert.",
     "3", "C", "Not Started"),
    ("E", "Failed signals + midtrend entries + (optional) pyramiding",
     "Operationalise Schwager's most important rule and the high-value midtrend entries.",
     "12 failed-signal detectors, reversal wiring, four midtrend entries, pyramiding with guards.",
     "F-040..F-043 green. In backtest on 5 tickers, failed-signal reversals improve net "
     "expectancy vs Phase-D baseline.",
     "3", "D", "Not Started"),
    ("F", "Backtest / walk-forward harness",
     "Validate the system honestly across multiple periods and markets with realistic costs.",
     "backtest_harness, walk-forward split, slippage model, parameter-set sensitivity report, "
     "BacktestStudio UI.",
     "Multi-period (3+) multi-market (5+) report with OOS expectancy, drawdown, win rate. No "
     "single 'optimum' relied on — variance across neighboring parameter sets reported.",
     "3", "E", "Not Started"),
    ("G", "Discipline overlay (Ch 16-17)",
     "Enforce the rules that make the system actually work in practice.",
     "discipline_checks (plan-before-trade, no-averaging-down, stop-widening, streaks), "
     "DisciplineDashboard, 82-rules advisory engine.",
     "Trader cannot save a stopless plan; averaging-down warned; drawdown-triggered size cut "
     "applied automatically; dashboard surfaces rule violations.",
     "2", "D", "Not Started"),
]
write_table(ws, headers, phases)
autosize(ws, [8, 50, 50, 60, 60, 12, 16, 14])
add_conditional(ws, "H", 2, 1 + len(phases), STATUS_FILLS)

# ---------------------------------------------------------------------------
# Sheet 6: Book_Chapter_Map
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Book_Chapter_Map")
headers = ["Chapter", "Title", "Part", "Core concepts (operational)",
           "Decision rules (extract)", "Computable artifacts (book code/)",
           "Maps to pipeline stage(s)"]

chapters = [
    ("3", "Trends", "1 — Basic tools",
     "Uptrend/downtrend via HH/HL pivots; trend lines / TD lines / internal trend lines / "
     "channels; MAs as smoothed trend proxy.",
     "Pullback entry on retest; trend-line penetration -> reversal alert; channel exceed -> "
     "profit-taking; MA cross alert.",
     "01_trend_pivots.py, 02_trend_lines.py, 03_td_lines.py, 04_internal_trend_line.py, "
     "05_moving_average.py",
     "s5_regime, s6L_direction, regime_inputs"),
    ("4", "Trading ranges and S/R", "1 — Basic tools",
     "Horizontal range bounded by S/R; concentration zones; envelopes (MA ± ATR); breakout "
     "termination of range.",
     "Long at support / short at resistance with stop just past; breakout confirmation; range "
     "break failure -> bull/bear trap.",
     "support_resistance.py, trading_range.py, concentration_zones.py, price_envelope.py, "
     "breakout.py",
     "s5_regime, s6R_range, s7R_breakout, s9R_failed"),
    ("5", "Chart patterns", "1 — Basic tools",
     "Reversal (H&S, double top/bottom, rounding); continuation (flag, pennant, triangle); "
     "gaps; bar patterns (WRD, dojis).",
     "Pattern completion -> measured-move target; pattern failure -> failed-signal reversal; "
     "pause-in-trend -> midtrend entry.",
     "chart_patterns.py",
     "s4_patterns, s10L_failed, s9R_failed, risk_targets"),
    ("6", "Oscillators", "1 — Basic tools",
     "Momentum, ROC, MA oscillator, MACD, RSI, Stochastics; divergence; overbought/oversold.",
     "TREND: oscillators are alert-only; require price confirm. RANGE: actionable at "
     "boundaries (osc + S/R).",
     "oscillators_master.py, generate_images.py",
     "s3_indicators, osc_note (overlay)"),
    ("7", "Is chart analysis still valid?", "1 — Basic tools",
     "Confirmation layers (penetration, time delay, pattern, oscillator agree, fundamentals); "
     "multi-timeframe filter; failed-signal exploitation; risk-control asymmetry.",
     "Need 2+ confirmations; short-term aligned with weekly/monthly; reverse on failed signal.",
     "ch7_master.py",
     "s9L_filter, s8R_confirm, s10L_failed"),
    ("8", "Midtrend entry & pyramiding", "2 — Trading issues",
     "% retracement (33/50/67); minor-reaction reversal; continuation breakout from flag; "
     "long-MA reaction.",
     "Add to winners only when unrealised P&L positive; max 50% from pyramids; tighter stops "
     "on add; no-add after N consecutive losses.",
     "01_midtrend_entry.py, 02_percent_retracement.py, 03_reversal_minor_reaction.py, "
     "04_continuation_breakout.py, 05_long_term_ma_reaction.py, 06_pyramiding.py",
     "s7L_signal (midtrend), state_machine (pyramid loop)"),
    ("9", "Stop-loss points", "2 — Trading issues",
     "Stop-loss liquidation point set BEFORE entry; GTC orders; trend-line / range / pattern / "
     "WRD / relative-high-low / money / trailing stop.",
     "Stop quality hierarchy: chart > swing > %/ATR > money. Stops never widen. Size = "
     "risk / (entry - stop).",
     "10 scripts: 01_stop_loss_liquidation_point.py .. 09_trailing_stop.py + 10_cic_anti_pattern.py",
     "risk_stops"),
    ("10", "Exit objectives", "2 — Trading issues",
     "Chart-based / measured-move / S/R / RSI-extreme / contrary-opinion objectives; trailing "
     "stop; change-of-market-opinion exit.",
     "Combine 2+ targets at same price -> high confidence; partial 50% at first target, trail "
     "the remainder.",
     "(inferred from index — measured-move calc, S/R objective, trailing executor)",
     "risk_targets"),
    ("11", "Failed signals (THE rule)", "2 — Trading issues",
     "12 failure types: bull/bear traps, false TL break, filled gaps, return to spike/WRD "
     "extremes, counter-anticipated flag, opposite breakout, H&S target failure, double-top "
     "break, rounding break, reliability tag.",
     "When pattern fails to deliver expected move -> reverse position; size smaller (still "
     "needs confirmation).",
     "generate_illustrations.py, chapter11_demo.py, chapter11_detectors.py",
     "s10L_failed, s9R_failed, failed_signal_principle"),
    ("12", "Real-world chart analysis", "2 — Trading issues",
     "19 actual trades; 12 cross-cutting concepts (premise violation, stop quality, measured "
     "move, traps, flags, counter-anticipated, WRD, trailing, island/spike, expectancy).",
     "Trade ends when premise is violated. Flexibility > prediction. Asymmetry near range "
     "extremes.",
     "chapter12_concepts.py, generate_charts.py, wavelet_decomposer.py, emd_stl_decomposer.py, "
     "chart_decomposer.py",
     "out_actions (illustrates every branch)"),
    ("13", "Charting / analysis software", "3 — Trading systems",
     "Data choices (historical / EOD / real-time); data cleaning; continuous futures rolls; "
     "software selection by timeframe.",
     "EOD sufficient for swing/position; clean before backtest; never trust uncleaned data.",
     "historical_price_data.py, end_of_day_data.py, real_time_data.py, software_selection.py, "
     "data_cleaning.py, continuous_futures.py",
     "s2_data"),
    ("14", "Trading systems: structure & design",
     "3 — Trading systems",
     "Fast vs slow MA trade-off; MA / breakout systems; confirmation (penetration, time, "
     "pattern); filters; market-char adjustments; buy/sell asymmetry; pyramiding; trade exit; "
     "countertrend overlay; 3-level diversification.",
     "Architecture template (8 steps); 10-point common-problems checklist.",
     "13 scripts: 01_trend_following_systems.py .. 13_diversification.py",
     "s7L_signal, s9L_filter, s8R_confirm, state_machine"),
    ("15", "Testing & optimising trading systems",
     "3 — Trading systems",
     "Well-chosen-example fallacy; parameter sets; time stability; slippage; overfitting; "
     "in-sample vs out-of-sample; walk-forward; multi-market.",
     "Optimize on train ONLY; evaluate on test; trade neighbors of optimum, not the optimum; "
     "test on 3+ markets.",
     "10 concept_*.py + 00_generate_illustrations.py",
     "backtest_harness"),
    ("16", "Planned trading approach", "4 — Practical guidelines",
     "Seven-step blueprint (philosophy, markets, risk, planning routine, notebook, diary, "
     "analysis). Position sizing = risk / (entry-stop); volatility / equity / streak "
     "adjustments.",
     "Risk per trade fixed before entry; reduce size after losing streak; segment trade "
     "analysis by direction / pattern / period.",
     "01_position_sizing.py, 02_volatility_sizing.py, 03_traders_notebook.py, "
     "04_trade_analysis.py",
     "risk_sizing, discipline, log_loop"),
    ("17", "82 trading rules", "4 — Practical guidelines",
     "82 rules across 7 sections: entries, exits/risk, money mgmt, holding winners, psychology, "
     "patterns, analysis/review.",
     "Rule 76: failed signal more reliable than original. Cut losers, hold winners. Plan "
     "doesn't bend in bull/bear. Bad habits never cured, only latent.",
     "7 scripts: 01_entering_trades.py .. 07_analysis_review.py",
     "discipline, all branches"),
]
write_table(ws, headers, chapters)
autosize(ws, [8, 32, 18, 60, 60, 56, 30])

# ---------------------------------------------------------------------------
# Sheet 7: Risks_Open_Questions
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Risks_Open_Questions")
headers = ["ID", "Category", "Question / Risk", "Why it matters",
           "Assumed answer (for now)", "Needed before phase", "Status"]

questions = [
    ("Q-01", "Market scope",
     "Are we targeting only Vietnamese stocks (HOSE/HNX/UPCOM) at first, or also crypto / US / "
     "futures?",
     "Liquidity, data source, regime parameters and pattern detectors all differ. Schwager's "
     "examples are futures/equities.",
     "DECIDED: VN equities only in v1 (HOSE/HNX/UPCOM via existing stock-data-service / vnstock).",
     "B", "Answered"),
    ("Q-02", "Bar frequency",
     "Daily bars only, or intraday too?",
     "Real-time / tick data adds infra cost (Ch 13) and changes stop/target logic.",
     "Daily bars only in v1.",
     "B", "Open"),
    ("Q-03", "Action mode",
     "Recommend only, paper trade, or eventually broker-integrated?",
     "Compliance, security, and the discipline layer differ. Auto-execute is far higher risk.",
     "DECIDED: Recommend-only in v1 (no execution, no broker). Paper-trade simulator may "
     "appear later via the Phase F backtest harness, but not as a v1 promise.",
     "B", "Answered"),
    ("Q-04", "Universe size",
     "How many tickers per user are we promising the helper supports per bar?",
     "Drives caching strategy, batch size in orchestrator, and DB sizing.",
     "Up to 30 tickers per user per bar in v1.",
     "B", "Open"),
    ("Q-05", "Long-only or long/short?",
     "Schwager treats both. VN cash equities are largely long-only retail.",
     "Halves complexity if long-only; SHORT state in state-machine can be deferred.",
     "DECIDED: Long + short both modelled. SHORT state shipped in v1. (Note: VN cash market "
     "doesn't allow true short — treat SHORT as 'short bias' / exit-then-wait recommendation, "
     "or paper-only.)",
     "C", "Answered"),
    ("Q-06", "Account / risk-pref source",
     "Where does account equity come from for sizing — user-typed, or portfolio-service "
     "valuation?",
     "Sizing is meaningless without a real equity figure.",
     "DECIDED: User-typed equity in v1 with manual override stored in user_risk_prefs. "
     "Portfolio-service valuation feed deferred to v2.",
     "C", "Answered"),
    ("Q-07", "Authoritative pattern detector",
     "Use book-Python scripts (knowledge/.../code/) directly, or re-implement inside services?",
     "Reuse saves time but couples our service to research code style.",
     "Re-implement inside ai-insights-service / decision-helper-service; book code as reference + test oracle.",
     "A", "Open"),
    ("Q-08", "Multi-timeframe inputs",
     "Do we require weekly + monthly bars for the multi-TF filter (Ch 7), or daily only?",
     "Materially changes data access pattern and orchestrator cost.",
     "Daily + weekly resampled from daily; no separate weekly fetch.",
     "B", "Open"),
    ("Q-09", "Failed-signal sensitivity",
     "Schwager's 12 detectors are sensitive to threshold tuning. Tune per-symbol, "
     "per-regime, or globally?",
     "Per-symbol = more accurate but harder to validate; global = simpler.",
     "Global defaults + per-user override; per-symbol learned in Phase F backtest.",
     "E", "Open"),
    ("Q-10", "Execution price assumption",
     "For backtest + paper-trade, do we use next-day open, close, or limit-at-signal?",
     "Changes backtest results materially.",
     "Next-day open in v1 (Schwager's typical assumption).",
     "F", "Open"),
    ("Q-11", "Discipline enforcement strength",
     "Does TradePlanModal HARD-block trades without stop+target, or just warn?",
     "Hard-block = strongest discipline (Ch 17 spirit) but most friction.",
     "DECIDED: Hard-block save when stop OR target is missing. Override flag exists but is "
     "recorded as a rule-violation in the journal.",
     "D", "Answered"),
    ("Q-12", "Prediction-service overlay",
     "Should ML predictions from prediction-service modify the recommendation, or stay "
     "advisory-only?",
     "ML overlay can amplify signal but also overfit. Schwager-style discipline is the "
     "primary edge.",
     "DECIDED: Advisory-only in v1. Surfaced in DecisionConsole sidebar; never used in "
     "stop / target / size / state transitions.",
     "B", "Answered"),
    ("Q-13", "Localisation / language",
     "Reason chain and discipline messages — English, Vietnamese, or both?",
     "Frontend i18n footprint.",
     "English in v1; Vietnamese in Phase G.",
     "D", "Open"),
    ("Q-14", "Auth model for helper service",
     "Reuse auth-service JWT, or new service-to-service token?",
     "Consistency with the rest of the stack vs additional risk surface.",
     "Reuse auth-service JWT (same as other services).",
     "B", "Open"),
    ("Q-15", "Decision cadence",
     "Run helper on every bar close, on user request, or both?",
     "Continuous run -> infra cost; on-request -> stale data risk.",
     "DECIDED: Daily scheduled run after market close for every watched ticker + on-demand "
     "recompute exposed via API (used by DecisionConsole 'Recompute now' button).",
     "C", "Answered"),
]
write_table(ws, headers, questions)
autosize(ws, [8, 22, 60, 60, 60, 18, 12])
add_conditional(ws, "G", 2, 1 + len(questions), STATUS_FILLS)

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
wb.save(OUT)
print(f"wrote {OUT}")
